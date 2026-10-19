"""
INDENT TOOL - single-file Streamlit app
=======================================

Works in 2 phases (they can be days apart):

Phase 1  Tab "1 · Build Indent Workbook"
    Download the input template, fill it, upload it, download TR_INDENT.xlsx (Sheet2).
    Input sheets    WH            HSN, Section, Barcode, Article, Color, Size, MRP, Qty
                    Sale          Store Code, Store Name, Section, Qty
                    Current Stock Store Code, Store Name, Section, Qty
                    Outward       Store Code, Store Name, Section, Qty   (optional)
    - Store Name may be empty in a store sheet; it is taken from the other store sheets.
    - Stores are ranked by total sale. Summary cells show sale/stock, or sale/stock/out when the
      Outward sheet has data. A store that is only in Outward is skipped with a warning.
    - Short instructions are written in the empty block at the top left (A1:H, rows grow with
      the number of sections).

Manual   Allocate quantities in the store columns of Sheet2 (column L onwards), save in Excel.

Phase 2  Tab "2 · Fill Platform Templates"
    Upload the allocated sheet, choose the indent date, download:
      FF Xpress Store Indent  FF_XPRESS_STORE_INDENT<DDMMYYYY>.xlsx
      PO Indent               PO_INDENT_<DD-Mon-YY>.xlsx
      Plain Data              INDENT_DATA_<DDMMYYYY>.xlsx  (all fields, to copy into any other format)

Files
    app.py                  this file (screen + all logic)
    requirements.txt        libraries, versions fixed so the app runs the same everywhere
    .streamlit/config.toml  light theme and upload limit
    templates/*.xlsx        blank platform templates (column mapping is in PLATFORMS below)

Run on a PC       pip install -r requirements.txt   then   streamlit run app.py
Streamlit Cloud   Put these files in a GitHub repo (private is fine). On share.streamlit.io choose
                  New app, pick the repo, main file app.py, Advanced settings: Python 3.13, Deploy.
                  Uploaded files are processed in memory only; nothing is saved.

Adding a fixed platform format (developer)
    Put the blank template in templates/ and add one entry to PLATFORMS. "columns" maps a template
    header (case and extra spaces ignored) to a field: store_code, store_name, ean, qty, mrp, hsn,
    section, article, color, size, indent_date, season_year - or {"const": "EA"} for a fixed value.
    "as": "int" or "text" is optional. Headers that are not listed stay blank.
"""
import datetime as dt
import io
import re
from copy import copy
from pathlib import Path

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# =================================================================== SETTINGS
TEMPLATE_DIR = Path(__file__).parent / "templates"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

WH_COLS = ["HSN", "Section", "Barcode", "Article", "Color", "Size", "MRP", "Qty"]
STORE_COLS = ["Store Code", "Store Name", "Section", "Qty"]
SHEETS = {"WH": WH_COLS, "Sale": STORE_COLS, "Current Stock": STORE_COLS, "Outward": STORE_COLS}
SIZE_ORDER = ["XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL"]

FONT = "Aptos Narrow"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
CENTER = Alignment(horizontal="center", vertical="center")
FIRST_STORE_COL = 12  # L
MIN_LAST_COL = 62     # BJ (formula range end, kept as in original)

# Written in the empty block A1:H<subtotal row> of the indent workbook.
INSTRUCTIONS = [
    "1. Enter the qty to send in each store column (from column L). Leave blank or 0 for no stock.",
    "2. Use whole numbers only. No decimals, negative numbers or text.",
    "3. Indent (column J) is the total allocated. If SHORT (column K) is negative, reduce the qty.",
    "4. The rows above show {label} for each store and section. Row 1 shows the store rank by sale.",
    "5. Do not delete or move the yellow header row, the store codes or columns A to K.",
    "6. Save the file in Excel, then upload it in tab 2 (Fill Platform Templates).",
]

# Platform output formats (Phase 2). One blank template per platform in templates/.
PLATFORMS = [
    {
        "name": "FF Xpress Store Indent",
        "template_file": "FF_XPRESS_STORE_INDENT.xlsx",
        "sheet": "FF Pune Handewadi",
        "header_row": 2,
        "first_data_row": 3,
        "output_name": "FF_XPRESS_STORE_INDENT{date:%d%m%Y}.xlsx",
        "columns": {
            "INDENT DATE": {"field": "indent_date"},
            "STORE CODE": {"field": "store_code"},
            "STORE NAME": {"field": "store_name"},
            "EAN": {"field": "ean", "as": "int"},
            "QTY": {"field": "qty"},
            "UNIT OF MEASURE": {"const": "N/A"},
            "MRP": {"field": "mrp"},
            "HSN": {"field": "hsn", "as": "int"},
            "SEASON YEAR": {"field": "season_year"},
            "PRODUCT": {"field": "section"},
            "Style NO/ Name": {"field": "article"},
            "Colour": {"field": "color"},
            "Size": {"field": "size"},
        },
    },
    {
        "name": "PO Indent",
        "template_file": "PO_INDENT_TEMPLATE.xlsx",
        "sheet": "IM_RECORDS",
        "header_row": 1,
        "first_data_row": 2,
        "output_name": "PO_INDENT_{date:%d-%b-%y}.xlsx",
        "columns": {
            "Store/DC Code": {"field": "store_code"},
            "Substore": {"field": "store_code"},
            "EAN": {"field": "ean", "as": "int"},
            "Total Qty": {"field": "qty"},
            "Unit Of Measure": {"const": "EA"},
            "MRP": {"field": "mrp"},
            "HSN": {"field": "hsn", "as": "int"},
            "Season Year": {"field": "season_year"},
        },
    },
]

# Plain list of everything that was allocated (no platform template).
PLAIN_DATA = {
    "name": "Plain Data",
    "output_name": "INDENT_DATA_{date:%d%m%Y}.xlsx",
    "columns": {
        "Indent Date": {"field": "indent_date"},
        "Store Code": {"field": "store_code"},
        "Store Name": {"field": "store_name"},
        "EAN": {"field": "ean", "as": "int"},
        "HSN": {"field": "hsn", "as": "int"},
        "Section": {"field": "section"},
        "Article": {"field": "article"},
        "Color": {"field": "color"},
        "Size": {"field": "size"},
        "MRP": {"field": "mrp"},
        "Qty": {"field": "qty"},
        "Season Year": {"field": "season_year"},
    },
    "widths": [11, 11, 34, 16, 11, 16, 12, 16, 7, 9, 7, 11],
}
OUTPUTS = PLATFORMS + [PLAIN_DATA]


# =================================================================== PHASE 1: input template
def make_template() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for i, _ in enumerate(cols, 1):
            c = ws.cell(1, i)
            c.font = Font(name=FONT, bold=True)
            c.fill = YELLOW
            c.alignment = CENTER
            c.border = BORDER
            ws.column_dimensions[get_column_letter(i)].width = 18
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# =================================================================== PHASE 1: reading
def _clean_text(s):
    s = s.where(s.notna(), "").astype(str).str.strip()
    return s.replace({"nan": "", "None": "", "<NA>": ""})


def _code(v):
    """Store code -> normalised text key (works for 1234, 1234.0, '1234', 'tbs1')."""
    if v is None or (isinstance(v, float) and v != v):
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).strip().upper()


def _store_cell(code):
    """Digits-only code (no leading zero) is written as a number, else as text."""
    if code.isdigit() and not code.startswith("0"):
        return int(code)
    return code


def _num(v):
    try:
        f = float(v)
        return int(f) if f == int(f) else f
    except (TypeError, ValueError):
        return v


def _read(file, sheet, cols, optional=False):
    """Read one input sheet. An optional sheet that is missing or has no rows returns None."""
    try:
        df = pd.read_excel(file, sheet_name=sheet, dtype=object)
    except ValueError:
        if optional:
            return None
        raise ValueError(f"Sheet '{sheet}' was not found in the file. Please use the sheet names from the input template.")
    df.columns = [str(c).strip() for c in df.columns]
    if optional and df.dropna(how="all").empty:
        return None
    missing = [c for c in cols if c not in df.columns]
    if missing:
        if sheet != "WH" and missing == ["Store Name"]:
            df["Store Name"] = ""
        else:
            raise ValueError(f"Sheet '{sheet}' is missing these columns: {', '.join(missing)}")
    df = df[cols].dropna(how="all").copy()
    return df


def _clean_store(df):
    df["Store Code"] = df["Store Code"].map(_code)
    df["Store Name"] = _clean_text(df["Store Name"])
    df["Section"] = _clean_text(df["Section"])
    df["Qty"] = pd.to_numeric(df["Qty"], errors="coerce").fillna(0)
    return df[df["Store Code"] != ""]


def load_inputs(file):
    """Returns (wh, sale, stock, out). out is None when the Outward sheet is missing or empty."""
    wh = _read(file, "WH", WH_COLS)
    sale = _clean_store(_read(file, "Sale", STORE_COLS))
    stock = _clean_store(_read(file, "Current Stock", STORE_COLS))
    out = _read(file, "Outward", STORE_COLS, optional=True)
    if out is not None:
        out = _clean_store(out)
        if out.empty:
            out = None
    for c in ("HSN", "Section", "Barcode", "Article", "Color", "Size"):
        wh[c] = wh[c].where(wh[c].notna(), "")
    wh["Section"] = _clean_text(wh["Section"])
    wh["Color"] = _clean_text(wh["Color"])
    wh["Size"] = _clean_text(wh["Size"]).str.upper()
    wh["MRP"] = pd.to_numeric(wh["MRP"], errors="coerce").fillna(0)
    wh["Qty"] = pd.to_numeric(wh["Qty"], errors="coerce").fillna(0)
    for c in ("HSN", "Barcode", "Article"):
        wh[c] = wh[c].map(lambda v: _num(v) if str(v).strip() != "" else "")
    return wh, sale, stock, out


# =================================================================== PHASE 1: calculation
def _fmt(v):
    return int(v) if float(v) == int(v) else round(float(v), 2)


def compute(wh, sale, stock, out=None):
    """Returns (wh_rows_df, stores_df, sections, sale_pivot, stock_pivot, out_pivot, warnings)."""
    warnings = []
    wh_g = (wh.groupby(["HSN", "Section", "Barcode", "Article", "Color", "Size", "MRP"],
                       dropna=False, sort=False, as_index=False)["Qty"].sum())

    def art_key(a):
        return (0, float(a), "") if isinstance(a, (int, float)) else (1, 0, str(a))

    def size_key(s):
        return SIZE_ORDER.index(s) if s in SIZE_ORDER else len(SIZE_ORDER)

    wh_g["_a"] = wh_g["Article"].map(art_key)
    wh_g["_c"] = wh_g["Color"].str.upper()
    wh_g["_s"] = wh_g["Size"].map(size_key)
    wh_g = wh_g.sort_values(["_a", "_c", "_s", "Size"], kind="stable").drop(columns=["_a", "_c", "_s"])
    odd = sorted(set(wh_g["Size"]) - set(SIZE_ORDER) - {""})
    if odd:
        warnings.append(f"These sizes are not in the standard list and were placed last: {', '.join(odd)}")

    # Outward is optional; a store that is only in Outward is skipped (warning at the end)
    out_only = []
    if out is not None:
        known = set(sale["Store Code"]) | set(stock["Store Code"])
        out_only = sorted(set(out["Store Code"]) - known)
        out = out[out["Store Code"].isin(known)].copy()
    store_dfs = (sale, stock) if out is None else (sale, stock, out)

    # sections: union, alphabetical (case-insensitive)
    secs = {}
    for df in (wh_g,) + store_dfs:
        for s in df["Section"]:
            if s and s.upper() not in secs:
                secs[s.upper()] = s
    sections = [secs[k] for k in sorted(secs)]
    wh_secs = {s.upper() for s in wh_g["Section"]}
    extra = [secs[k] for k in secs if k not in wh_secs]
    if extra:
        src = "Sale/Stock" if out is None else "Sale/Stock/Outward"
        warnings.append(f"These sections exist in {src} but not in WH (they are included in the summary rows): {', '.join(extra)}")

    for df in store_dfs:
        df["_sk"] = df["Section"].str.upper()
    sp = sale.groupby(["Store Code", "_sk"])["Qty"].sum()
    kp = stock.groupby(["Store Code", "_sk"])["Qty"].sum()
    op = None if out is None else out.groupby(["Store Code", "_sk"])["Qty"].sum()

    # store names: first non-empty from Sale, then Stock, then Outward
    names = {}
    for df in store_dfs:
        for code, nm in zip(df["Store Code"], df["Store Name"]):
            if nm and code not in names:
                names[code] = nm
    codes = sorted(set(sale["Store Code"]) | set(stock["Store Code"]))
    noname = [c for c in codes if c not in names]
    if noname:
        warnings.append(f"Store name not found for these stores (the code is shown instead): {', '.join(noname)}")

    tot_sale = sale.groupby("Store Code")["Qty"].sum()
    tot_stock = stock.groupby("Store Code")["Qty"].sum()
    stores = pd.DataFrame({"code": codes})
    stores["name"] = stores["code"].map(lambda c: names.get(c, c))
    stores["sale"] = stores["code"].map(lambda c: tot_sale.get(c, 0))
    stores["stock"] = stores["code"].map(lambda c: tot_stock.get(c, 0))
    if out is not None:
        tot_out = out.groupby("Store Code")["Qty"].sum()
        stores["out"] = stores["code"].map(lambda c: tot_out.get(c, 0))
    stores = stores.sort_values(["sale", "stock", "code"], ascending=[False, False, True],
                                kind="stable").reset_index(drop=True)
    stores["rank"] = range(1, len(stores) + 1)
    if out_only:
        warnings.append(f"These stores are only in Outward (not in Sale or Current Stock) and were skipped: {', '.join(out_only)}")
    return wh_g, stores, sections, sp, kp, op, warnings


# =================================================================== PHASE 1: workbook
def _write_instructions(ws, last_row, label):
    """Instructions in the empty block A1:H<last_row>. The block grows with the number of sections."""
    lines = [s.format(label=label) for s in INSTRUCTIONS]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    head = ws.cell(1, 1, "INSTRUCTIONS")
    head.font = Font(name=FONT, size=10, bold=True, color="FFFF0000")
    head.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=last_row, end_column=8)
    body = ws.cell(2, 1, "\n".join(lines))
    body.font = Font(name=FONT, size=9)
    body.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    rows = last_row - 1
    need = 12.5 * len(lines) + 3          # points needed for the text
    if 15 * rows < need:                  # few sections: make these rows a little taller
        h = round(need / rows + 0.5)
        for r in range(2, last_row + 1):
            ws.row_dimensions[r].height = h


def build_workbook(wh_g, stores, sections, sp, kp, op=None):
    n = len(sections)
    r_gt = 2
    r_sub = 3 + n
    r_name = 4 + n
    r_hdr = 5 + n
    r_data = 6 + n
    ns = len(stores)
    last_store_col = FIRST_STORE_COL + max(ns, 1) - 1
    last_col = max(last_store_col, 11)
    sum_end = get_column_letter(max(MIN_LAST_COL, last_store_col))

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet2"

    def put(r, c, v=None, size=11, bold=False, fill=None, border=True, color=None, wrap=False, fmt=None, center=True):
        cell = ws.cell(r, c, v)
        cell.font = Font(name=FONT, size=size, bold=bold, color=color)
        if center:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        if fill:
            cell.fill = fill
        if border:
            cell.border = BORDER
        if fmt:
            cell.number_format = fmt
        return cell

    # row 1
    for c in (9, 10, 11):
        put(1, c, "Sale Wise Store Ranking" if c == 9 else None, color="FFFF0000")
    ws.merge_cells(start_row=1, start_column=9, end_row=1, end_column=11)
    for i, row in stores.iterrows():
        put(1, FIRST_STORE_COL + i, int(row["rank"]), color="FFFF0000", center=False)

    # rows 2..(2+n): Sale/INV, or Sale/INV/OUT with Outward (only Grand Total row is yellow)
    def join(s, k, o):
        t = f"{_fmt(s)}/{_fmt(k)}"
        return t if op is None else f"{t}/{_fmt(o)}"

    def cell_text(row, skey):
        if skey is None:
            return join(row["sale"], row["stock"], row["out"] if op is not None else 0)
        key = (row["code"], skey)
        return join(sp.get(key, 0), kp.get(key, 0), op.get(key, 0) if op is not None else 0)

    def total_text(skey):
        if skey is None:
            return join(stores["sale"].sum(), stores["stock"].sum(),
                        stores["out"].sum() if op is not None else 0)
        sale_t = sum(v for (c, k), v in sp.items() if k == skey)
        stock_t = sum(v for (c, k), v in kp.items() if k == skey)
        out_t = sum(v for (c, k), v in op.items() if k == skey) if op is not None else 0
        return join(sale_t, stock_t, out_t)

    labels = ["Grand Total"] + sections
    keys = [None] + [s.upper() for s in sections]
    for k, (lab, key) in enumerate(zip(labels, keys)):
        r = r_gt + k
        fill = YELLOW if k == 0 else None
        put(r, 9, None, size=8)
        put(r, 10, total_text(key), size=8, fill=fill)
        put(r, 11, lab, size=8, fill=fill)
        for i, row in stores.iterrows():
            put(r, FIRST_STORE_COL + i, cell_text(row, key), size=8, fill=fill)

    # subtotal row
    for c in range(9, last_col + 1):
        L = get_column_letter(c)
        put(r_sub, c, f"=SUBTOTAL(9,{L}{r_data}:{L}111365)")

    # store names row
    put(r_name, 1, "WH STOCK", size=36, bold=True, color="FFFF0000")
    for c in range(2, 9):
        put(r_name, c)
    ws.merge_cells(start_row=r_name, start_column=1, end_row=r_name, end_column=8)
    put(r_name, 9, "Calculation", size=26, bold=True, color="FFFF0000")
    put(r_name, 10)
    put(r_name, 11)
    ws.merge_cells(start_row=r_name, start_column=9, end_row=r_name, end_column=11)
    for i, row in stores.iterrows():
        put(r_name, FIRST_STORE_COL + i, row["name"], size=8, bold=True, wrap=True)
    ws.row_dimensions[r_name].height = 56.25

    # header
    hdr = ["HSN", "Section", "Barcode", "Article", "SKU", "Color", "Size", "MRP", " Qty", "Indent", "SHORT"]
    for c, h in enumerate(hdr, 1):
        put(r_hdr, c, h, fill=YELLOW)
    for i, row in stores.iterrows():
        put(r_hdr, FIRST_STORE_COL + i, _store_cell(row["code"]), fill=YELLOW)

    # data
    for k, row in enumerate(wh_g.itertuples(index=False)):
        r = r_data + k
        qty = _fmt(row.Qty)
        vals = [row.HSN, row.Section, row.Barcode, row.Article, f"=D{r}&F{r}", row.Color,
                row.Size, _fmt(row.MRP), qty, f"=SUM(L{r}:{sum_end}{r})", f"=I{r}-J{r}"]
        for c, v in enumerate(vals, 1):
            put(r, c, v if v != "" else None, fmt="0" if c in (3, 5) else None)

    widths = {"A": 11.14, "B": 14.14, "C": 17.0, "E": 24.43, "F": 19.0, "G": 7.57, "K": 12.71}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    if op is not None:  # three numbers per cell: widen J and store columns only when needed
        for c in [10] + list(range(FIRST_STORE_COL, FIRST_STORE_COL + ns)):
            longest = max(len(str(ws.cell(r, c).value or "")) for r in range(r_gt, r_gt + len(labels)))
            need = round(longest * 0.75 + 1.5, 2)
            if need > 8.43:
                ws.column_dimensions[get_column_letter(c)].width = need

    _write_instructions(ws, r_sub, "Sale/Stock" if op is None else "Sale/Stock/Outward")

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def run(file):
    wh, sale, stock, out = load_inputs(file)
    if wh.empty:
        raise ValueError("The WH sheet is empty.")
    wh_g, stores, sections, sp, kp, op, warnings = compute(wh, sale, stock, out)
    if stores.empty:
        raise ValueError("No stores were found in the Sale or Current Stock sheets.")
    data = build_workbook(wh_g, stores, sections, sp, kp, op)
    summary = {"skus": len(wh_g), "stores": len(stores), "sections": sections,
               "wh_qty": _fmt(wh_g["Qty"].sum()), "outward": op is not None}
    return data, stores, summary, warnings


# =================================================================== PHASE 2: read allocation
def _norm(s):
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip().casefold()


def _whole(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def read_allocation(file):
    """Returns (records, store_summary, totals, warnings)."""
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb["Sheet2"] if "Sheet2" in wb.sheetnames else wb.worksheets[0]

    hdr = None
    for r in range(1, min(ws.max_row, 60) + 1):
        if _norm(ws.cell(r, 1).value) == "hsn" and _norm(ws.cell(r, 2).value) == "section":
            hdr = r
            break
    if hdr is None:
        raise ValueError("Header row (HSN, Section, Barcode...) was not found. Please upload the sheet generated in Step 1.")
    name_row = hdr - 1

    stores, seen, warnings = [], set(), []
    for c in range(FIRST_STORE_COL, ws.max_column + 1):
        code = ws.cell(hdr, c).value
        if _blank(code):
            continue
        code = _whole(code)
        if isinstance(code, str):
            code = code.strip()
        if str(code) in seen:
            warnings.append(f"Store code {code} appears more than once in the header (column {get_column_letter(c)}).")
        seen.add(str(code))
        nm = ws.cell(name_row, c).value
        stores.append((c, code, "" if nm is None else str(nm).strip()))
    if not stores:
        raise ValueError("No store columns were found (expected from column L onwards).")

    per_store = {c: [] for c, _, _ in stores}
    bad, over, nobar = [], [], []
    for r in range(hdr + 1, ws.max_row + 1):
        row_vals = {}
        total = 0
        for c, code, nm in stores:
            v = ws.cell(r, c).value
            if _blank(v):
                continue
            if isinstance(v, str):
                try:
                    v = float(v.strip())
                except ValueError:
                    bad.append(ws.cell(r, c).coordinate)
                    continue
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                bad.append(ws.cell(r, c).coordinate)
                continue
            if v == 0:
                continue
            if v < 0 or v != int(v):
                bad.append(ws.cell(r, c).coordinate)
                continue
            row_vals[c] = int(v)
            total += int(v)
        if not row_vals:
            continue

        meta = {
            "hsn": _whole(ws.cell(r, 1).value),
            "section": ws.cell(r, 2).value,
            "ean": _whole(ws.cell(r, 3).value),
            "article": _whole(ws.cell(r, 4).value),
            "color": ws.cell(r, 6).value,
            "size": ws.cell(r, 7).value,
            "mrp": _whole(ws.cell(r, 8).value),
        }
        if _blank(meta["ean"]):
            nobar.append(r)
        wh_qty = ws.cell(r, 9).value
        if isinstance(wh_qty, (int, float)) and total > wh_qty:
            over.append(f"Row {r} (Article {meta['article']}, {meta['color']}, {meta['size']}): allocated {total} > WH Qty {_whole(wh_qty)}")
        for c, q in row_vals.items():
            per_store[c].append({**meta, "qty": q})

    records, store_rows = [], []
    for c, code, nm in stores:
        recs = per_store[c]
        for rec in recs:
            rec["store_code"] = code
            rec["store_name"] = nm
        records.extend(recs)
        store_rows.append({"Code": code, "Store": nm, "Rows": len(recs), "Qty": sum(x["qty"] for x in recs)})

    if bad:
        warnings.append(f"{len(bad)} cell(s) had an invalid quantity (text / negative / decimal) and were skipped: "
                        + ", ".join(bad[:10]) + (" ..." if len(bad) > 10 else ""))
    if over:
        warnings.append(f"{len(over)} SKU(s) are allocated more than the WH Qty (negative SHORT):\n" + "\n".join(over[:10])
                        + ("\n..." if len(over) > 10 else ""))
    if nobar:
        warnings.append(f"{len(nobar)} allocated row(s) have an empty Barcode (rows: "
                        + ", ".join(map(str, nobar[:10])) + (" ..." if len(nobar) > 10 else "") + ")")
    totals = {"stores": sum(1 for s in store_rows if s["Rows"]), "rows": len(records),
              "qty": sum(x["qty"] for x in records)}
    if not records:
        raise ValueError("No allocation (quantity) was found in any store column.")
    return records, store_rows, totals, warnings


# =================================================================== PHASE 2: outputs
def _resolve(spec, rec, indent_date):
    if "const" in spec:
        v = spec["const"]
    else:
        f = spec["field"]
        if f == "indent_date":
            v = dt.datetime.combine(indent_date, dt.time())
        elif f == "season_year":
            v = indent_date.year
        else:
            v = rec.get(f)
    if _blank(v):
        return None
    kind = spec.get("as")
    if kind == "int":
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return v
    if kind == "text":
        return str(_whole(v))
    return v


def fill_template(cfg, records, indent_date):
    """Fill a platform template. Returns (file_name, bytes, warnings)."""
    wb = openpyxl.load_workbook(TEMPLATE_DIR / cfg["template_file"])
    ws = wb[cfg["sheet"]] if cfg["sheet"] in wb.sheetnames else wb.worksheets[0]
    hdr_row, first = cfg["header_row"], cfg["first_data_row"]

    heads = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(hdr_row, c).value
        if v is not None:
            heads[_norm(v)] = c
    colmap, warnings = [], []
    for name, spec in cfg["columns"].items():
        c = heads.get(_norm(name))
        if c is None:
            warnings.append(f"{cfg['name']}: column '{name}' was not found in the template and was skipped.")
            continue
        colmap.append((c, spec))

    max_r, max_c = ws.max_row, ws.max_column
    proto = {c: copy(ws.cell(first, c)._style) for c in range(1, max_c + 1)}
    for r in range(first, max_r + 1):          # remove sample rows, keep formatting
        for c in range(1, max_c + 1):
            ws.cell(r, c).value = None

    for i, rec in enumerate(records):
        r = first + i
        if r > max_r:
            for c in range(1, max_c + 1):
                ws.cell(r, c)._style = copy(proto[c])
        for c, spec in colmap:
            v = _resolve(spec, rec, indent_date)
            if v is None:
                continue
            cell = ws.cell(r, c)
            cell.value = v
            if spec.get("field") == "indent_date" and cell.number_format == "General":
                cell.number_format = "d-mmm-yy"

    bio = io.BytesIO()
    wb.save(bio)
    return cfg["output_name"].format(date=indent_date), bio.getvalue(), warnings


def plain_data(cfg, records, indent_date):
    """All allocated rows with every field, no platform template. Returns (file_name, bytes, warnings)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Indent Data"
    heads = list(cfg["columns"])
    ws.append(heads)
    for i, w in enumerate(cfg["widths"], 1):
        c = ws.cell(1, i)
        c.font = Font(name=FONT, bold=True)
        c.fill = YELLOW
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, rec in enumerate(records, 2):
        for i, spec in enumerate(cfg["columns"].values(), 1):
            ws.cell(r, i, _resolve(spec, rec, indent_date))
    for r in range(2, len(records) + 2):
        ws.cell(r, heads.index("Indent Date") + 1).number_format = "d-mmm-yy"
        ws.cell(r, heads.index("EAN") + 1).number_format = "0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(records) + 1}"
    bio = io.BytesIO()
    wb.save(bio)
    return cfg["output_name"].format(date=indent_date), bio.getvalue(), []


def make_output(cfg, records, indent_date):
    if "template_file" in cfg:
        return fill_template(cfg, records, indent_date)
    return plain_data(cfg, records, indent_date)


# =================================================================== SCREEN
st.set_page_config(page_title="Indent Tool", layout="wide", initial_sidebar_state="collapsed")

st.markdown(
    """
    <style>
    /* ---- layout: use the full window ---- */
    .block-container {max-width: 100% !important; padding: 1.1rem 2.6rem 1.2rem 2.6rem;}
    #MainMenu, footer {visibility: hidden;}
    header[data-testid="stHeader"] {background: transparent;}

    /* ---- hero banner ---- */
    .hero {background: linear-gradient(115deg, #0B1B4D 0%, #1D4ED8 60%, #3B82F6 100%);
           border-radius: 16px; padding: 20px 28px; margin-bottom: 14px; color: #fff;
           display: flex; align-items: center; justify-content: space-between;
           box-shadow: 0 6px 18px rgba(29,78,216,.22);}
    .hero-title {font-size: 1.75rem; font-weight: 750; letter-spacing: .2px; line-height: 1.15;}
    .hero-sub {font-size: .98rem; opacity: .88; margin-top: 4px;}
    .hero-flow {display: flex; gap: 10px; align-items: center; font-size: .85rem;}
    .pill {background: rgba(255,255,255,.16); border: 1px solid rgba(255,255,255,.28);
           border-radius: 999px; padding: 6px 14px; font-weight: 600; white-space: nowrap;}
    .arrow {opacity: .7;}

    /* ---- tabs ---- */
    .stTabs [data-baseweb="tab-list"] {gap: 6px; border-bottom: 1px solid #DDE3EE;}
    .stTabs [data-baseweb="tab"] {height: 48px; padding: 0 22px; font-weight: 650; font-size: 1rem;
                                  border-radius: 10px 10px 0 0;}
    .stTabs [aria-selected="true"] {background: #fff; color: #1D4ED8;}

    /* ---- panels ---- */
    div[data-testid="stVerticalBlockBorderWrapper"] {background: #fff; border-radius: 14px;
                                                    border-color: #E3E8F1 !important;}
    .panel-title {font-size: 1.05rem; font-weight: 700; color: #0F172A; margin-bottom: 2px;}
    .panel-sub {font-size: .85rem; color: #64748B; margin-bottom: 10px;}
    .step {display: flex; align-items: center; gap: 10px; font-weight: 650; margin: 14px 0 6px 0;
           font-size: .98rem;}
    .step .num {background: #1D4ED8; color: #fff; width: 24px; height: 24px; border-radius: 50%;
                display: inline-flex; align-items: center; justify-content: center; font-size: .8rem;}
    .hint {font-size: .82rem; color: #64748B; margin: 2px 0 8px 34px;}
    .empty {border: 1.5px dashed #CBD5E1; border-radius: 12px; padding: 54px 20px; text-align: center;
            color: #64748B; background: #FAFBFD;}
    .empty b {color: #334155; font-size: 1.02rem;}

    /* ---- metrics ---- */
    div[data-testid="stMetric"] {background: #F6F8FC; border: 1px solid #E3E8F1; border-radius: 12px;
                                 padding: 12px 16px;}
    div[data-testid="stMetricLabel"] p {font-size: .82rem; color: #64748B; font-weight: 600;}
    div[data-testid="stMetricValue"] {font-size: 1.65rem; font-weight: 750; color: #0F172A;}

    /* ---- buttons: full width, rounded ---- */
    .stButton > button, .stDownloadButton > button {width: 100%; border-radius: 10px; font-weight: 650;
                                                    padding: .62rem 1rem;}
    div[data-testid="stFileUploader"] section {border-radius: 12px;}
    </style>
    <div class="hero">
      <div>
        <div class="hero-title">Indent Tool</div>
        <div class="hero-sub">Build the indent workbook, allocate stock to stores, then generate platform upload templates.</div>
      </div>
      <div class="hero-flow">
        <span class="pill">WH + Sale + Stock (+ Outward)</span><span class="arrow">&rarr;</span>
        <span class="pill">Indent workbook</span><span class="arrow">&rarr;</span>
        <span class="pill">Allocate in Excel</span><span class="arrow">&rarr;</span>
        <span class="pill">Platform templates</span>
      </div>
    </div>
    """,
    unsafe_allow_html=True,
)


def step(n, text, hint=None):
    st.markdown(f'<div class="step"><span class="num">{n}</span>{text}</div>', unsafe_allow_html=True)
    if hint:
        st.markdown(f'<div class="hint">{hint}</div>', unsafe_allow_html=True)


def empty(title, text):
    st.markdown(f'<div class="empty"><b>{title}</b><br>{text}</div>', unsafe_allow_html=True)


def table(df, height):
    """st.dataframe; a column mixing numbers and text (store code 1234 and TBS1) is shown as text."""
    df = df.copy()
    for c in df.columns:
        if df[c].dtype == object and df[c].dropna().map(type).nunique() > 1:
            df[c] = df[c].map(lambda v: v if v is None else str(v))
    st.dataframe(df, hide_index=True, height=height)


tab1, tab2 = st.tabs(["1 · Build Indent Workbook", "2 · Fill Platform Templates"])

# ------------------------------------------------------------------ TAB 1
with tab1:
    left, right = st.columns([5, 8], gap="large")

    with left:
        with st.container(border=True):
            st.markdown('<div class="panel-title">Inputs</div>'
                        '<div class="panel-sub">One Excel file with the sheets WH, Sale, Current Stock '
                        'and Outward (optional).</div>',
                        unsafe_allow_html=True)
            step(1, "Download the input template",
                 "Fill the WH, Sale and Current Stock sheets. Outward is optional: fill it to see "
                 "sale/stock/out, or leave it empty. The Store Name column may be left empty; it is "
                 "looked up from the other store sheets.")
            st.download_button("Download input template", make_template(),
                               file_name="Indent_Input_Template.xlsx", mime=XLSX, key="dl_tpl")
            step(2, "Upload the filled template")
            up1 = st.file_uploader("Filled input template", type=["xlsx"], key="up1",
                                   label_visibility="collapsed")
            step(3, "Generate")
            go1 = st.button("Build indent workbook", type="primary", key="go1", disabled=up1 is None)
            if go1:
                try:
                    with st.spinner("Building workbook..."):
                        st.session_state["r1"] = (up1.name, run(up1))
                except Exception as e:
                    st.session_state.pop("r1", None)
                    st.error(str(e))

    with right:
        with st.container(border=True):
            st.markdown('<div class="panel-title">Result</div>', unsafe_allow_html=True)
            if "r1" not in st.session_state:
                empty("No workbook yet",
                      "Upload the filled input template and click <i>Build indent workbook</i>.<br>"
                      "The store ranking and the download will appear here.")
            else:
                fname, (data, stores, summ, warns) = st.session_state["r1"]
                note = " · Outward included" if summ["outward"] else ""
                st.markdown(f'<div class="panel-sub">Generated from {fname}{note}</div>', unsafe_allow_html=True)
                m1, m2, m3, m4 = st.columns(4)
                m1.metric("WH SKUs", f"{summ['skus']:,}")
                m2.metric("WH Qty", f"{summ['wh_qty']:,}")
                m3.metric("Stores", summ["stores"])
                m4.metric("Sections", len(summ["sections"]))
                for w in warns:
                    st.warning(w)
                st.markdown('<div class="step" style="margin-top:12px">Store ranking (by sale, descending)</div>',
                            unsafe_allow_html=True)
                cols = ["Rank", "Code", "Store", "Sale", "Stock"] + (["Outward"] if summ["outward"] else [])
                view = stores.rename(columns={"rank": "Rank", "code": "Code", "name": "Store",
                                              "sale": "Sale", "stock": "Stock", "out": "Outward"})[cols]
                table(view, min(38 * (len(view) + 1) + 4, 330))
                st.download_button("Download indent workbook (Sheet2)", data, file_name="TR_INDENT.xlsx",
                                   mime=XLSX, type="primary", key="dl1")

# ------------------------------------------------------------------ TAB 2
with tab2:
    names = [t["name"] for t in OUTPUTS]
    left, right = st.columns([5, 8], gap="large")

    with left:
        with st.container(border=True):
            st.markdown('<div class="panel-title">Inputs</div>'
                        '<div class="panel-sub">Upload the indent sheet after allocating quantities to the stores.</div>',
                        unsafe_allow_html=True)
            step(1, "Upload the allocated indent sheet",
                 "Enter the quantities in the store columns of Sheet2 and save the file in Excel before uploading.")
            up2 = st.file_uploader("Allocated indent sheet", type=["xlsx"], key="up2",
                                   label_visibility="collapsed")
            step(2, "Indent date and templates",
                 "Plain Data has every field in one simple sheet, to copy into any other format.")
            indent_date = st.date_input("Indent date", value=dt.date.today(), format="DD/MM/YYYY")
            chosen = st.multiselect("Templates to generate", names, default=names)
            step(3, "Generate")
            go2 = st.button("Fill templates", type="primary", key="go2",
                            disabled=(up2 is None or not chosen))
            if go2:
                try:
                    with st.spinner("Filling templates..."):
                        records, store_rows, totals, warns = read_allocation(up2)
                        files, fwarns = [], []
                        for t in OUTPUTS:
                            if t["name"] in chosen:
                                fn, blob, w = make_output(t, records, indent_date)
                                files.append((t["name"], fn, blob))
                                fwarns += w
                    st.session_state["r2"] = (up2.name, records, files, store_rows, totals, warns + fwarns)
                except Exception as e:
                    st.session_state.pop("r2", None)
                    st.error(str(e))

    with right:
        with st.container(border=True):
            st.markdown('<div class="panel-title">Result</div>', unsafe_allow_html=True)
            if "r2" not in st.session_state:
                empty("No templates yet",
                      "Upload the allocated indent sheet and click <i>Fill templates</i>.<br>"
                      "The store-wise summary, a preview and the downloads will appear here.")
            else:
                fname, records, files, store_rows, totals, warns = st.session_state["r2"]
                st.markdown(f'<div class="panel-sub">Generated from {fname}</div>', unsafe_allow_html=True)
                m1, m2, m3, m4 = st.columns(4)
                m1.metric("Stores allocated", totals["stores"])
                m2.metric("Rows", f"{totals['rows']:,}")
                m3.metric("Total Qty", f"{totals['qty']:,}")
                m4.metric("Templates", len(files))
                for w in warns:
                    st.warning(w)

                sc, pc = st.columns([4, 7], gap="medium")
                with sc:
                    st.markdown('<div class="step" style="margin-top:12px">Store-wise summary</div>',
                                unsafe_allow_html=True)
                    sdf = pd.DataFrame(store_rows)
                    table(sdf, min(38 * (len(sdf) + 1) + 4, 330))
                with pc:
                    st.markdown('<div class="step" style="margin-top:12px">Preview (first 100 rows)</div>',
                                unsafe_allow_html=True)
                    pv = pd.DataFrame(records[:100]).rename(columns={
                        "store_code": "Store Code", "store_name": "Store", "ean": "EAN", "qty": "Qty",
                        "mrp": "MRP", "hsn": "HSN", "section": "Section", "article": "Article",
                        "color": "Color", "size": "Size"})
                    cols = ["Store Code", "Store", "EAN", "Section", "Article", "Color", "Size", "MRP", "Qty"]
                    table(pv[cols], 330)

                st.markdown('<div class="step" style="margin-top:12px">Downloads</div>', unsafe_allow_html=True)
                dcols = st.columns(len(files))
                for i, (name, fn, blob) in enumerate(files):
                    dcols[i].download_button(f"Download {name}", blob, file_name=fn, mime=XLSX,
                                             type="primary", key=f"dl2_{i}")
                    dcols[i].caption(fn)
